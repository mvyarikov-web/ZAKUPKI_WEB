"""Blueprint для страниц (index, view)."""
import os
from flask import Blueprint, render_template, jsonify, request, current_app
from markupsafe import Markup
from app.services.files import allowed_file
from app.services.state import FilesState
from app.services.indexing import get_index_path

pages_bp = Blueprint('pages', __name__)


def _get_files_state():
    """Получить экземпляр FilesState для текущего приложения."""
    results_file = current_app.config['SEARCH_RESULTS_FILE']
    return FilesState(results_file)


@pages_bp.route('/')
def index():
    """Главная страница"""
    # Загружаем сохраненные результаты поиска
    files_state = _get_files_state()
    last_search_terms = files_state.get_last_search_terms()
    file_status = files_state.get_file_status()
    
    # Получаем список загруженных файлов с группировкой по папкам
    files_by_folder = {}
    total_files = 0
    
    current_app.logger.info(f"Проверяем папку: {current_app.config['UPLOAD_FOLDER']}")
    
    if os.path.exists(current_app.config['UPLOAD_FOLDER']):
        # Рекурсивно обходим все файлы и папки
        for root, dirs, files in os.walk(current_app.config['UPLOAD_FOLDER']):
            for filename in files:
                # Скрываем служебный индексный файл и временные файлы Office
                if filename == '_search_index.txt' or filename.startswith('~$') or filename.startswith('$'):
                    continue
                file_path = os.path.join(root, filename)
                # Определяем относительную папку
                relative_folder = os.path.relpath(root, current_app.config['UPLOAD_FOLDER'])

                # Правильно определяем название папки
                if relative_folder == '.':
                    folder_display_name = '📁 Загруженные файлы'
                    folder_key = 'root'
                else:
                    folder_parts = relative_folder.split(os.sep)
                    original_folder_name = folder_parts[-1]
                    folder_display_name = f'📂 {original_folder_name}'
                    folder_key = relative_folder

                file_size = os.path.getsize(file_path)
                file_key = os.path.join(relative_folder, filename) if relative_folder != '.' else filename
                file_data = file_status.get(file_key, {})
                status = file_data.get('status', 'not_checked')
                # Если формат не поддерживается — пометим явно
                if not allowed_file(filename, current_app.config['ALLOWED_EXTENSIONS']):
                    status = 'unsupported'
                    file_data = {**file_data, 'status': 'unsupported', 'error': 'Неподдерживаемый формат'}
                    files_state.set_file_status(file_key, 'unsupported', 
                                               {'error': 'Неподдерживаемый формат'})

                display_name = file_data.get('original_name', filename)
                file_info = {
                    'name': display_name,
                    'size': file_size,
                    'status': status,
                    'path': file_key,
                    'relative_folder': relative_folder
                }

                if folder_key not in files_by_folder:
                    files_by_folder[folder_key] = {
                        'display_name': folder_display_name,
                        'relative_path': relative_folder,
                        'files': []
                    }

                files_by_folder[folder_key]['files'].append(file_info)
                total_files += 1
                current_app.logger.debug(f"Добавлен файл: {filename} в папку {folder_display_name}, размер: {file_size}, статус: {status}")
    else:
        current_app.logger.warning("Папка uploads не существует")
    
    current_app.logger.info(f"Всего файлов для отображения: {total_files}, папок: {len(files_by_folder)}")
    return render_template(
        'index.html',
        files_by_folder=files_by_folder,
        total_files=total_files,
        last_search_terms=last_search_terms,
        file_status=file_status,
    )


@pages_bp.get('/view/<path:filepath>')
def view_file(filepath: str):
    """Отображение содержимого файла с подсветкой ключевых слов."""
    from urllib.parse import unquote
    from app.services.files import is_safe_subpath, allowed_file
    
    # Импорты для извлечения текста
    import docx
    import pdfplumber
    import openpyxl
    import docx2txt
    import subprocess
    
    try:
        decoded_filepath = unquote(filepath)
        current_app.logger.info(f"Просмотр файла: {decoded_filepath}")
        
        # Проверка безопасности пути
        if not is_safe_subpath(current_app.config['UPLOAD_FOLDER'], decoded_filepath):
            return jsonify({'error': 'Недопустимый путь к файлу'}), 400
        
        file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], decoded_filepath)
        
        if not os.path.exists(file_path) or not os.path.isfile(file_path):
            return jsonify({'error': 'Файл не найден'}), 404
        
        # Проверяем статус файла
        files_state = _get_files_state()
        file_data = files_state.get_file_status(decoded_filepath)
        
        status = file_data.get('status', 'not_checked')
        char_count = file_data.get('char_count')
        error = file_data.get('error')
        
        # Запрещаем чтение неподдерживаемых, ошибочных и пустых файлов
        if status == 'unsupported':
            return jsonify({'error': 'Файл неподдерживаемого формата'}), 403
        if status == 'error' or error:
            return jsonify({'error': f'Ошибка обработки файла: {error}'}), 403
        if char_count == 0:
            return jsonify({'error': 'Файл пуст или не содержит текста'}), 403
        
        # Извлекаем текст в зависимости от расширения
        ext = os.path.splitext(file_path)[1].lower()
        text = ''
        
        # Быстрые форматы
        if ext in ('.txt', '.csv', '.tsv', '.xml', '.json', '.html', '.htm'):
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        elif ext == '.pdf':
            try:
                with pdfplumber.open(file_path) as pdf:
                    text = '\n'.join(page.extract_text() or '' for page in pdf.pages[:50])
            except Exception:
                current_app.logger.exception('Ошибка извлечения текста из PDF')
        elif ext == '.docx':
            try:
                text = docx2txt.process(file_path)
            except Exception:
                try:
                    doc = docx.Document(file_path)
                    text = '\n'.join(p.text for p in doc.paragraphs)
                except Exception:
                    current_app.logger.exception('Ошибка извлечения текста из DOCX')
        elif ext == '.doc':
            try:
                result = subprocess.run(
                    ['antiword', file_path],
                    capture_output=True, text=True, timeout=10
                )
                if result.returncode == 0:
                    text = result.stdout
            except Exception:
                current_app.logger.exception('Ошибка извлечения текста из DOC')
        elif ext in ('.xlsx', '.xls'):
            try:
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                parts = []
                for sheet in wb.worksheets[:5]:
                    for row in sheet.iter_rows(max_row=100, values_only=True):
                        parts.append(' '.join(str(c) for c in row if c))
                text = '\n'.join(parts)
            except Exception:
                current_app.logger.exception('Ошибка извлечения текста из XLSX/XLS')
        
        if not text:
            return jsonify({'error': 'Не удалось извлечь текст из файла'}), 403
        
        # Получаем ключевые слова из query параметра
        query = request.args.get('q', '')
        keywords = [k.strip() for k in query.split(',') if k.strip()] if query else []
        
        # Подсветка ключевых слов
        if keywords:
            import re
            for keyword in keywords:
                pattern = re.compile(re.escape(keyword), re.IGNORECASE)
                text = pattern.sub(lambda m: f'<mark>{m.group(0)}</mark>', text)
        
        # Экранируем HTML для безопасности (кроме наших mark)
        # text уже содержит <mark>, поэтому используем Markup
        safe_text = Markup(text.replace('<', '&lt;').replace('>', '&gt;')
                          .replace('&lt;mark&gt;', '<mark>')
                          .replace('&lt;/mark&gt;', '</mark>'))
        
        return render_template('view.html',
                             filename=os.path.basename(decoded_filepath),
                             content=safe_text,
                             keywords=keywords)
    
    except Exception as e:
        current_app.logger.exception('view_file error')
        return jsonify({'error': str(e)}), 500
